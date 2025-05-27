import tkinter as tk
from tkinter import ttk, colorchooser, messagebox, filedialog
import PyPDF2
from docx import Document
import os

ACCENT = "#00FFC6"
BG = "#181A1B"
FG = "#F8F8F2"
BTN_BG = "#232526"
BTN_FG = "#00FFC6"
FONT_MAIN = ("Segoe UI", 48, "bold")
FONT_BTN = ("Segoe UI", 13, "bold")
FONT_TEXT = ("Segoe UI", 13)

class SwiftReadApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Swift Read")
        self.words = []
        self.word_index = 0
        self.is_reading = False
        self.theme_bg = BG
        self.theme_fg = FG
        self.accent = ACCENT
        self.current_after = None

        self.root.configure(bg=self.theme_bg)
        self.root.geometry("900x600")  # Increased height for new controls
        self.root.resizable(False, False)

        style = ttk.Style()
        style.theme_use("clam")
        style.configure("Accent.TButton", font=FONT_BTN, background=BTN_BG, foreground=BTN_FG, borderwidth=0)
        style.map("Accent.TButton",
                  background=[("active", self.accent)],
                  foreground=[("active", BG)])

        self.display = tk.Label(root, text="Swift Read", font=FONT_MAIN, bg=self.theme_bg, fg=self.accent, pady=30)
        self.display.pack(fill="x", pady=(30, 10))

        # File upload frame
        file_frame = tk.Frame(root, bg=self.theme_bg)
        file_frame.pack(pady=5)
        
        tk.Label(file_frame, text="Upload a file:", bg=self.theme_bg, fg=self.theme_fg, font=FONT_TEXT).grid(row=0, column=0, padx=5)
        ttk.Button(file_frame, text="PDF", style="Accent.TButton", command=lambda: self.upload_file("pdf")).grid(row=0, column=1, padx=5)
        ttk.Button(file_frame, text="Word", style="Accent.TButton", command=lambda: self.upload_file("word")).grid(row=0, column=2, padx=5)
        ttk.Button(file_frame, text="Text", style="Accent.TButton", command=lambda: self.upload_file("txt")).grid(row=0, column=3, padx=5)

        text_frame = tk.Frame(root, bg=self.theme_bg)
        text_frame.pack(pady=5)
        tk.Label(text_frame, text="Or paste/type your text below:", bg=self.theme_bg, fg=self.theme_fg, font=FONT_TEXT).pack(anchor="w")
        self.textbox = tk.Text(text_frame, height=5, width=80, font=FONT_TEXT, bg=BTN_BG, fg=FG, insertbackground=FG, wrap="word", borderwidth=2, relief="flat")
        self.textbox.pack(pady=5)
        
        btn_frame = tk.Frame(text_frame, bg=self.theme_bg)
        btn_frame.pack(fill="x")
        ttk.Button(btn_frame, text="Load Text", style="Accent.TButton", command=self.load_text).pack(side="right", pady=2)
        ttk.Button(btn_frame, text="Clear", style="Accent.TButton", command=self.clear_text).pack(side="right", padx=5, pady=2)

        controls = tk.Frame(root, bg=self.theme_bg)
        controls.pack(pady=10)

        ttk.Button(controls, text="Theme", style="Accent.TButton", command=self.change_theme).grid(row=0, column=0, padx=8)
        ttk.Button(controls, text="Start", style="Accent.TButton", command=self.start_reading).grid(row=0, column=1, padx=8)
        ttk.Button(controls, text="Pause", style="Accent.TButton", command=self.pause_reading).grid(row=0, column=2, padx=8)
        ttk.Button(controls, text="Reset", style="Accent.TButton", command=self.reset).grid(row=0, column=3, padx=8)

        slider_frame = tk.Frame(root, bg=self.theme_bg)
        slider_frame.pack(pady=10)
        tk.Label(slider_frame, text="Speed:", bg=self.theme_bg, fg=self.theme_fg, font=("Segoe UI", 12)).pack(side="left")
        self.speed_var = tk.IntVar(value=300)
        self.speed_slider = ttk.Scale(slider_frame, from_=100, to=1000, orient="horizontal", variable=self.speed_var, length=300, command=self.update_speed_label)
        self.speed_slider.pack(side="left", padx=10)
        self.speed_label = tk.Label(slider_frame, text=f"{self.speed_var.get()} WPM", bg=self.theme_bg, fg=self.accent, font=("Segoe UI", 12, "bold"))
        self.speed_label.pack(side="left")

        self.progress_label = tk.Label(root, text="", bg=self.theme_bg, fg=self.theme_fg, font=("Segoe UI", 11))
        self.progress_label.pack(pady=5)

        self.file_info_label = tk.Label(root, text="", bg=self.theme_bg, fg=self.accent, font=("Segoe UI", 10))
        self.file_info_label.pack(pady=5)

    def update_speed_label(self, event=None):
        self.speed_label.config(text=f"{self.speed_var.get()} WPM")

    def upload_file(self, file_type):
        filetypes = []
        if file_type == "pdf":
            filetypes = [("PDF files", "*.pdf")]
        elif file_type == "word":
            filetypes = [("Word files", "*.docx"), ("Word files", "*.doc")]
        elif file_type == "txt":
            filetypes = [("Text files", "*.txt")]
        
        file_path = filedialog.askopenfilename(filetypes=filetypes)
        if not file_path:
            return
        
        try:
            text = ""
            if file_type == "pdf":
                text = self.extract_text_from_pdf(file_path)
            elif file_type == "word":
                text = self.extract_text_from_word(file_path)
            elif file_type == "txt":
                with open(file_path, 'r', encoding='utf-8') as f:
                    text = f.read()
            
            self.textbox.delete("1.0", "end")
            self.textbox.insert("1.0", text)
            self.file_info_label.config(text=f"Loaded: {os.path.basename(file_path)}")
            messagebox.showinfo("Success", "File content loaded successfully!")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to read file:\n{str(e)}")

    def extract_text_from_pdf(self, file_path):
        text = ""
        with open(file_path, 'rb') as file:
            reader = PyPDF2.PdfReader(file)
            for page in reader.pages:
                text += page.extract_text()
        return text

    def extract_text_from_word(self, file_path):
        doc = Document(file_path)
        return "\n".join([para.text for para in doc.paragraphs])

    def load_text(self):
        text = self.textbox.get("1.0", "end").strip()
        if not text:
            messagebox.showinfo("No Text", "Please paste/type some text or upload a file.")
            return
        self.words = text.split()
        self.word_index = 0
        self.display.config(text="Ready!", fg=self.accent)
        self.progress_label.config(text=f"{len(self.words)} words loaded.")

    def clear_text(self):
        self.textbox.delete("1.0", "end")
        self.words = []
        self.word_index = 0
        self.display.config(text="Swift Read", fg=self.accent)
        self.progress_label.config(text="")
        self.file_info_label.config(text="")

    def start_reading(self):
        if not self.words:
            self.load_text()
            if not self.words:
                return
        self.is_reading = True
        if self.current_after:
            self.root.after_cancel(self.current_after)
        self.show_next_word()

    def show_next_word(self):
        if not self.is_reading:
            return
        if self.word_index < len(self.words):
            word = self.words[self.word_index]
            self.display.config(text=word, fg=self.accent)
            self.progress_label.config(text=f"Word {self.word_index+1}/{len(self.words)}")
            self.word_index += 1
            delay = int(60000 / self.speed_var.get())
            self.current_after = self.root.after(delay, self.show_next_word)
        else:
            self.display.config(text="Done!", fg=self.accent)
            self.progress_label.config(text="All done!")
            self.is_reading = False

    def pause_reading(self):
        self.is_reading = False
        if self.current_after:
            self.root.after_cancel(self.current_after)
            self.current_after = None

    def reset(self):
        self.pause_reading()
        self.word_index = 0
        self.display.config(text="Swift Read", fg=self.accent)
        self.progress_label.config(text="")

    def change_theme(self):
        bg_color = colorchooser.askcolor(title="Choose background color", initialcolor=self.theme_bg)[1]
        fg_color = colorchooser.askcolor(title="Choose text color", initialcolor=self.theme_fg)[1]
        accent_color = colorchooser.askcolor(title="Choose accent color", initialcolor=self.accent)[1]
        if bg_color and fg_color and accent_color:
            self.theme_bg = bg_color
            self.theme_fg = fg_color
            self.accent = accent_color
            self.root.configure(bg=self.theme_bg)
            self.display.config(bg=self.theme_bg, fg=self.accent)
            self.textbox.config(bg=BTN_BG, fg=self.theme_fg, insertbackground=self.theme_fg)
            self.progress_label.config(bg=self.theme_bg, fg=self.theme_fg)
            self.file_info_label.config(bg=self.theme_bg, fg=self.accent)

if __name__ == "__main__":
    root = tk.Tk()
    
    # Add this function to handle icon paths for both development and EXE
    def resource_path(relative_path):
        try:
            base_path = sys._MEIPASS
        except Exception:
            base_path = os.path.abspath(".")
        return os.path.join(base_path, relative_path)
    
    # Set the icon
    try:
        root.iconbitmap(resource_path('s.ico'))
    except:
        # Fallback if icon fails to load
        pass
        
    app = SwiftReadApp(root)
    root.mainloop()